import sys
import os
import pptx
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

def create_deck():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_slide_layout = prs.slide_layouts[6] # blank layout

    # Colors
    c_navy = RGBColor(15, 44, 89)      # #0F2C59
    c_subnavy = RGBColor(26, 54, 93)   # #1A365D
    c_blue_accent = RGBColor(37, 99, 235) # #2563EB
    c_teal_accent = RGBColor(13, 148, 136) # #0D9488
    c_bg = RGBColor(248, 250, 252)     # #F8FAFC
    c_card_bg = RGBColor(255, 255, 255)
    c_border = RGBColor(226, 232, 240) # #E2E8F0
    c_text_dark = RGBColor(15, 23, 42) # #0F172A
    c_text_muted = RGBColor(71, 85, 105) # #475569
    c_white = RGBColor(255, 255, 255)
    c_gold = RGBColor(217, 119, 6)     # #D97706

    def set_bg(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = c_bg
        bg.line.fill.background()

    def add_header(slide, title_text, slide_num):
        # Header banner
        hdr = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(1.1))
        hdr.fill.solid()
        hdr.fill.fore_color.rgb = c_navy
        hdr.line.fill.background()

        # Title text
        tx_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.15), Inches(8.5), Inches(0.8))
        tf = tx_box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title_text
        p.font.name = "Calibri"
        p.font.size = Pt(26)
        p.font.bold = True
        p.font.color.rgb = c_white

        # SIH Subtitle inside header
        p2 = tf.add_paragraph()
        p2.text = "SMART INDIA HACKATHON 2026  |  PROBLEM STATEMENT ID: SIH26123"
        p2.font.name = "Calibri"
        p2.font.size = Pt(11)
        p2.font.color.rgb = RGBColor(203, 213, 225)

        # Team Badge Right
        badge = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(10.2), Inches(0.25), Inches(2.3), Inches(0.6))
        badge.fill.solid()
        badge.fill.fore_color.rgb = c_teal_accent
        badge.line.fill.background()
        tf_b = badge.text_frame
        p_b = tf_b.paragraphs[0]
        p_b.text = f"TEAM: SPARK-08  |  {slide_num}"
        p_b.alignment = PP_ALIGN.CENTER
        p_b.font.name = "Calibri"
        p_b.font.size = Pt(13)
        p_b.font.bold = True
        p_b.font.color.rgb = c_white

    def add_card(slide, left, top, width, height, title=None):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = c_card_bg
        card.line.color.rgb = c_border
        card.line.width = Pt(1)

        if title:
            # Card Header Title Bar inside
            tb = slide.shapes.add_textbox(left + Inches(0.15), top + Inches(0.1), width - Inches(0.3), Inches(0.4))
            tf = tb.text_frame
            p = tf.paragraphs[0]
            p.text = title
            p.font.name = "Calibri"
            p.font.size = Pt(16)
            p.font.bold = True
            p.font.color.rgb = c_subnavy

        return card

    # ==========================================
    # SLIDE 1: TITLE PAGE
    # ==========================================
    s1 = prs.slides.add_slide(blank_slide_layout)
    set_bg(s1)

    # Top Accent Bar
    bar = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(0.3))
    bar.fill.solid()
    bar.fill.fore_color.rgb = c_gold
    bar.line.fill.background()

    # Main Hero Title Box
    hero = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(0.6), Inches(11.733), Inches(2.4))
    hero.fill.solid()
    hero.fill.fore_color.rgb = c_navy
    hero.line.fill.background()

    tf = hero.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "SMART INDIA HACKATHON 2026"
    p.font.name = "Calibri"
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = c_gold

    p_sub = tf.add_paragraph()
    p_sub.text = "Edge-AI Based Distributed Fleet Coordination for Autonomous Mobile Robots (AMRs) in Smart Warehouses"
    p_sub.font.name = "Calibri"
    p_sub.font.size = Pt(20)
    p_sub.font.bold = True
    p_sub.font.color.rgb = c_white

    # Metadata Grid Cards
    # Card 1: Problem Details
    add_card(s1, Inches(0.8), Inches(3.2), Inches(5.7), Inches(3.8), "PROBLEM STATEMENT DETAILS")
    tb1 = s1.shapes.add_textbox(Inches(1.0), Inches(3.8), Inches(5.3), Inches(3.0))
    tf1 = tb1.text_frame
    tf1.word_wrap = True

    items1 = [
        ("• Problem Statement ID:", " SIH26123"),
        ("• Problem Statement Title:", " Edge-AI Based Distributed Fleet Coordination for Autonomous Mobile Robots (AMRs) in Smart Warehouses"),
        ("• Theme:", " Smart Automation"),
        ("• PS Category:", " Software"),
        ("• Organization / Nodal Agency:", " Bharat Electronics Limited")
    ]
    for label, val in items1:
        p = tf1.add_paragraph()
        run1 = p.add_run()
        run1.text = label
        run1.font.bold = True
        run1.font.size = Pt(13)
        run1.font.color.rgb = c_subnavy
        run2 = p.add_run()
        run2.text = val
        run2.font.size = Pt(13)
        run2.font.color.rgb = c_text_dark

    # Card 2: Team Details
    add_card(s1, Inches(6.833), Inches(3.2), Inches(5.7), Inches(3.8), "TEAM & SUBMISSION DETAILS")
    tb2 = s1.shapes.add_textbox(Inches(7.033), Inches(3.8), Inches(5.3), Inches(3.0))
    tf2 = tb2.text_frame
    tf2.word_wrap = True

    items2 = [
        ("• Team Name:", " SPARK-08"),
        ("• Team ID:", " 143472"),
        ("• Application Name:", " NEXUS-Fleet Digital Twin"),
        ("• Target Architecture:", " P2P Distributed Edge-AI Fleet Coordination"),
        ("• Submission Stage:", " SIH 2026 Grand Finale Solution Deck")
    ]
    for label, val in items2:
        p = tf2.add_paragraph()
        run1 = p.add_run()
        run1.text = label
        run1.font.bold = True
        run1.font.size = Pt(13)
        run1.font.color.rgb = c_subnavy
        run2 = p.add_run()
        run2.text = val
        run2.font.size = Pt(13)
        run2.font.color.rgb = c_text_dark


    # ==========================================
    # SLIDE 2: PROPOSED SOLUTION
    # ==========================================
    s2 = prs.slides.add_slide(blank_slide_layout)
    set_bg(s2)
    add_header(s2, "PROPOSED SOLUTION", 2)

    # Top Left: How It Addresses the Problem
    add_card(s2, Inches(0.8), Inches(1.3), Inches(5.7), Inches(3.4), "How It addresses the Problem")
    tb_address = s2.shapes.add_textbox(Inches(0.95), Inches(1.8), Inches(5.4), Inches(2.8))
    tf = tb_address.text_frame
    tf.word_wrap = True
    addr_points = [
        ("Decentralized Communication: ", "AMRs exchange pose, intent, task and reservation state directly; no central fleet brain makes every decision."),
        ("Dynamic Conflict Resolution: ", "Robots predict route/ETA conflicts at intersections and narrow aisles, then negotiate who yields or reroutes."),
        ("Task Allocation & Re-routing: ", "Robots independently bid for tasks and re-plan when an aisle becomes blocked or another robot fails."),
        ("Edge Execution: ", "The RobotAgent is designed to run locally on Raspberry Pi / Jetson-class hardware."),
        ("Fleet Visibility: ", "The existing 3D dashboard observes live positions, battery, tasks, reservations and events.")
    ]
    for title, desc in addr_points:
        p = tf.add_paragraph()
        run1 = p.add_run()
        run1.text = "• " + title
        run1.font.bold = True
        run1.font.size = Pt(11)
        run1.font.color.rgb = c_subnavy
        run2 = p.add_run()
        run2.text = desc
        run2.font.size = Pt(11)
        run2.font.color.rgb = c_text_dark

    # Top Right: Detailed Explanation
    add_card(s2, Inches(6.833), Inches(1.3), Inches(5.7), Inches(3.4), "Detailed explanation of the proposed solution")
    tb_det = s2.shapes.add_textbox(Inches(6.983), Inches(1.8), Inches(5.4), Inches(2.8))
    tf_det = tb_det.text_frame
    tf_det.word_wrap = True
    p_det = tf_det.paragraphs[0]
    p_det.text = "NEXUS-Fleet turns each AMR into an autonomous edge agent. Every robot maintains a local world model, plans its own route, evaluates tasks, communicates intent with peers, requests time-based corridor reservations, detects deadlocks and recovers from blocked aisles or communication loss. The web application remains the digital-twin and observability layer while the same agent architecture is prepared for physical edge deployment."
    p_det.font.size = Pt(12)
    p_det.font.color.rgb = c_text_dark

    # Bottom Left: Autonomous Coordination Workflow (5 Cards Across)
    add_card(s2, Inches(0.8), Inches(4.85), Inches(7.5), Inches(2.3), "Autonomous Coordination Workflow")
    steps = [
        ("1", "OBSERVE", "pose • peers • tasks"),
        ("2", "PREDICT", "ETA • conflict • risk"),
        ("3", "NEGOTIATE", "priority • reservation"),
        ("4", "ACT", "route • safe motion"),
        ("5", "RECOVER", "reroute • reassign")
    ]
    card_w = Inches(1.35)
    start_x = Inches(0.95)
    for idx, (num, label, desc) in enumerate(steps):
        cx = start_x + idx * Inches(1.42)
        scard = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, cx, Inches(5.35), card_w, Inches(1.6))
        scard.fill.solid()
        scard.fill.fore_color.rgb = c_navy if idx % 2 == 0 else c_subnavy
        scard.line.fill.background()
        stf = scard.text_frame
        stf.word_wrap = True
        p1 = stf.paragraphs[0]
        p1.text = num
        p1.font.bold = True
        p1.font.size = Pt(14)
        p1.font.color.rgb = c_gold
        p1.alignment = PP_ALIGN.CENTER
        p2 = stf.add_paragraph()
        p2.text = label
        p2.font.bold = True
        p2.font.size = Pt(11)
        p2.font.color.rgb = c_white
        p2.alignment = PP_ALIGN.CENTER
        p3 = stf.add_paragraph()
        p3.text = desc
        p3.font.size = Pt(9)
        p3.font.color.rgb = RGBColor(226, 232, 240)
        p3.alignment = PP_ALIGN.CENTER

    # Bottom Right: Innovation & Uniqueness
    add_card(s2, Inches(8.5), Inches(4.85), Inches(4.033), Inches(2.3), "Innovation & Uniqueness")
    tb_inn = s2.shapes.add_textbox(Inches(8.65), Inches(5.3), Inches(3.733), Inches(1.75))
    tf_inn = tb_inn.text_frame
    tf_inn.word_wrap = True
    inn_items = [
        ("Robot-as-Agent: ", "fleet autonomy is distributed across robots rather than hidden inside a central controller."),
        ("Spatio-temporal reservations: ", "a corridor/intersection is reserved for a robot over a time window, not merely marked occupied."),
        ("Explainable autonomy: ", "each yield, reroute, task award and recovery records a structured reason."),
        ("Measured evidence: ", "stop-and-wait is benchmarked under identical scenarios instead of hard-coding an improvement claim.")
    ]
    for title, desc in inn_items:
        p = tf_inn.add_paragraph()
        run1 = p.add_run()
        run1.text = "• " + title
        run1.font.bold = True
        run1.font.size = Pt(9.5)
        run1.font.color.rgb = c_subnavy
        run2 = p.add_run()
        run2.text = desc
        run2.font.size = Pt(9.5)
        run2.font.color.rgb = c_text_dark


    # ==========================================
    # SLIDE 3: TECHNICAL APPROACH
    # ==========================================
    s3 = prs.slides.add_slide(blank_slide_layout)
    set_bg(s3)
    add_header(s3, "TECHNICAL APPROACH", 3)

    # 3 Column Layout
    # Col 1: Technologies to be used
    add_card(s3, Inches(0.8), Inches(1.3), Inches(3.7), Inches(5.8), "Technologies to be used")
    tb_tech = s3.shapes.add_textbox(Inches(0.95), Inches(1.8), Inches(3.4), Inches(5.2))
    tf_tech = tb_tech.text_frame
    tf_tech.word_wrap = True
    tech_items = [
        ("Frontend: ", "React / Next.js + TypeScript + existing Three.js warehouse scene."),
        ("Backend: ", "Python, FastAPI, asyncio and WebSockets."),
        ("Agent Runtime: ", "Independent RobotAgent processes with local state and decision policy."),
        ("Planning: ", "Warehouse graph + A* route planning with congestion, battery and reservation costs."),
        ("Coordination: ", "Peer messages, task bidding, reservations, negotiation and deadlock recovery."),
        ("Safety: ", "Local collision monitoring, stale-peer handling and safe stop/slow behavior."),
        ("Storage: ", "PostgreSQL / event ledger for tasks, events, reservations and benchmark runs.")
    ]
    for label, val in tech_items:
        p = tf_tech.add_paragraph()
        r1 = p.add_run()
        r1.text = "• " + label
        r1.font.bold = True
        r1.font.size = Pt(11)
        r1.font.color.rgb = c_subnavy
        r2 = p.add_run()
        r2.text = val
        r2.font.size = Pt(10.5)
        r2.font.color.rgb = c_text_dark

    # Col 2: Methodology & Process (Flow Cards)
    add_card(s3, Inches(4.8), Inches(1.3), Inches(3.7), Inches(5.8), "Methodology & Implementation Process")
    flow_steps = [
        ("EXISTING 3D DIGITAL TWIN", "renders live robot state"),
        ("FASTAPI + WEBSOCKET", "real-time state / events"),
        ("ROBOT AGENT RUNTIME", "independent decision loops"),
        ("P2P COORDINATION", "intent • bids • reservations"),
        ("PLANNING + SAFETY", "A* • conflict • deadlock • safe motion")
    ]
    y_start = Inches(1.8)
    for idx, (title, sub) in enumerate(flow_steps):
        fc = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(4.95), y_start + idx * Inches(0.98), Inches(3.4), Inches(0.72))
        fc.fill.solid()
        fc.fill.fore_color.rgb = c_navy if idx % 2 == 0 else c_subnavy
        fc.line.fill.background()
        ftf = fc.text_frame
        ftf.word_wrap = True
        p1 = ftf.paragraphs[0]
        p1.text = title
        p1.font.bold = True
        p1.font.size = Pt(11)
        p1.font.color.rgb = c_gold
        p1.alignment = PP_ALIGN.CENTER
        p2 = ftf.add_paragraph()
        p2.text = sub
        p2.font.size = Pt(9.5)
        p2.font.color.rgb = c_white
        p2.alignment = PP_ALIGN.CENTER

        if idx < 4:
            # Arrow indicator
            arrow = s3.shapes.add_textbox(Inches(4.95), y_start + idx * Inches(0.98) + Inches(0.7), Inches(3.4), Inches(0.3))
            ap = arrow.text_frame.paragraphs[0]
            ap.text = "↓"
            ap.font.bold = True
            ap.font.size = Pt(14)
            ap.font.color.rgb = c_teal_accent
            ap.alignment = PP_ALIGN.CENTER

    # Col 3: Core Algorithms
    add_card(s3, Inches(8.8), Inches(1.3), Inches(3.733), Inches(5.8), "Core Algorithms")
    tb_alg = s3.shapes.add_textbox(Inches(8.95), Inches(1.8), Inches(3.433), Inches(5.2))
    tf_alg = tb_alg.text_frame
    tf_alg.word_wrap = True
    alg_items = [
        ("Task bidding: ", "ETA + distance + battery + workload + conflict risk."),
        ("A*: ", "dynamic route cost over the warehouse graph."),
        ("Conflict prediction: ", "route/edge/node/ETA overlap before collision."),
        ("Reservations: ", "time-bounded corridor/intersection leases."),
        ("Negotiation: ", "priority, wait time, urgency and fairness."),
        ("Deadlock: ", "wait-for graph cycle detection and recovery."),
        ("Rerouting: ", "blocked-edge invalidation and alternate path planning."),
        ("Benchmark: ", "completion time, waiting, distance, collisions, deadlocks and throughput.")
    ]
    for label, val in alg_items:
        p = tf_alg.add_paragraph()
        r1 = p.add_run()
        r1.text = "• " + label
        r1.font.bold = True
        r1.font.size = Pt(10.5)
        r1.font.color.rgb = c_subnavy
        r2 = p.add_run()
        r2.text = val
        r2.font.size = Pt(10)
        r2.font.color.rgb = c_text_dark


    # ==========================================
    # SLIDE 4: FEASIBILITY AND VIABILITY
    # ==========================================
    s4 = prs.slides.add_slide(blank_slide_layout)
    set_bg(s4)
    add_header(s4, "FEASIBILITY AND VIABILITY", 4)

    # 2 Big Columns
    # Left: Feasibility
    add_card(s4, Inches(0.8), Inches(1.3), Inches(5.7), Inches(4.5), "Feasibility")
    tb_feas = s4.shapes.add_textbox(Inches(0.95), Inches(1.8), Inches(5.4), Inches(3.8))
    tf_feas = tb_feas.text_frame
    tf_feas.word_wrap = True
    feas_items = [
        ("Simulation-first: ", "the expected solution calls for a multi-robot simulation before physical deployment."),
        ("Existing UI: ", "the current Three.js warehouse is retained, reducing implementation risk on visualization."),
        ("Edge-ready: ", "RobotAgent logic is separated from the browser and can be packaged for local edge hardware."),
        ("Computationally practical: ", "A*, graph reasoning, reservations and wait-for graphs are lightweight compared with cloud-scale inference."),
        ("Incremental MVP: ", "3 AMRs → coordination → failure scenarios → benchmark → edge deployment path.")
    ]
    for label, val in feas_items:
        p = tf_feas.add_paragraph()
        r1 = p.add_run()
        r1.text = "• " + label
        r1.font.bold = True
        r1.font.size = Pt(12)
        r1.font.color.rgb = c_subnavy
        r2 = p.add_run()
        r2.text = val
        r2.font.size = Pt(11.5)
        r2.font.color.rgb = c_text_dark

    # Right: Viability
    add_card(s4, Inches(6.833), Inches(1.3), Inches(5.7), Inches(4.5), "Viability")
    tb_viab = s4.shapes.add_textbox(Inches(6.983), Inches(1.8), Inches(5.4), Inches(3.8))
    tf_viab = tb_viab.text_frame
    tf_viab.word_wrap = True
    viab_items = [
        ("Warehouse productivity: ", "reduces unnecessary waiting by coordinating shared-space access."),
        ("Resilience: ", "blocked aisles, robot failures and stale communication become recoverable states."),
        ("Explainability: ", "event ledger records why an agent selected a task, yielded or rerouted."),
        ("Scalability: ", "typed APIs and peer messages support more robots without redesigning the UI."),
        ("Evidence-based evaluation: ", "same map, tasks, seed and initial conditions are used for baseline comparison.")
    ]
    for label, val in viab_items:
        p = tf_viab.add_paragraph()
        r1 = p.add_run()
        r1.text = "• " + label
        r1.font.bold = True
        r1.font.size = Pt(12)
        r1.font.color.rgb = c_subnavy
        r2 = p.add_run()
        r2.text = val
        r2.font.size = Pt(11.5)
        r2.font.color.rgb = c_text_dark

    # Bottom Full-Width Card: Key Validation
    add_card(s4, Inches(0.8), Inches(6.0), Inches(11.733), Inches(1.1), "Key validation")
    tb_val = s4.shapes.add_textbox(Inches(0.95), Inches(6.4), Inches(11.433), Inches(0.6))
    tf_val = tb_val.text_frame
    tf_val.word_wrap = True
    p_val = tf_val.paragraphs[0]
    p_val.text = "Key validation: zero inter-robot collisions is a test criterion; ≥20% task-time reduction is a target to be measured, not a pre-claimed result."
    p_val.font.size = Pt(13)
    p_val.font.bold = True
    p_val.font.color.rgb = c_teal_accent


    # ==========================================
    # SLIDE 5: IMPACT AND BENEFITS
    # ==========================================
    s5 = prs.slides.add_slide(blank_slide_layout)
    set_bg(s5)
    add_header(s5, "IMPACT AND BENEFITS", 5)

    # Left Side: Potential impact on the target audience (8 Points Grid)
    add_card(s5, Inches(0.8), Inches(1.3), Inches(6.8), Inches(5.8), "Potential impact on the target audience")
    tb_imp = s5.shapes.add_textbox(Inches(0.95), Inches(1.8), Inches(6.5), Inches(5.2))
    tf_imp = tb_imp.text_frame
    tf_imp.word_wrap = True
    impact_items = [
        ("1. Safer shared-space operation", "Predict conflicts before robots physically meet and apply local yielding / rerouting."),
        ("2. Reduced waiting", "Time-based reservations replace unnecessary stop-and-wait at recurring choke points."),
        ("3. Faster task completion", "Benchmark makes completion-time improvement measurable under overlapping paths."),
        ("4. Faster response to blockage", "Blocked aisles trigger route invalidation and autonomous replanning."),
        ("5. Resilient fleet operation", "A robot can fail or lose communication without freezing the whole fleet."),
        ("6. Explainable decisions", "Operators can inspect the exact reason for task, reservation and recovery decisions."),
        ("7. Edge deployment path", "Local RobotAgent execution reduces dependence on continuous cloud connectivity."),
        ("8. Scalable warehouse coordination", "The same architecture can support additional AMRs and scenarios.")
    ]
    for title, desc in impact_items:
        p = tf_imp.add_paragraph()
        r1 = p.add_run()
        r1.text = title + " — "
        r1.font.bold = True
        r1.font.size = Pt(10.5)
        r1.font.color.rgb = c_subnavy
        r2 = p.add_run()
        r2.text = desc
        r2.font.size = Pt(10)
        r2.font.color.rgb = c_text_dark

    # Right Side Top: Benefits of the solution (6 Grid Cards)
    add_card(s5, Inches(7.8), Inches(1.3), Inches(4.733), Inches(4.3), "Benefits of the solution")
    b_cards = [
        ("DECENTRALIZED", "No single fleet decision-maker"),
        ("PREDICTIVE", "Conflict detection before collision"),
        ("ADAPTIVE", "Blocked-aisle rerouting"),
        ("RESILIENT", "Communication / failure recovery"),
        ("EXPLAINABLE", "Event ledger + decision reasons"),
        ("MEASURABLE", "Baseline benchmark + KPIs")
    ]
    b_start_x = Inches(7.95)
    b_start_y = Inches(1.8)
    for idx, (title, desc) in enumerate(b_cards):
        col = idx % 2
        row = idx // 2
        bx = b_start_x + col * Inches(2.25)
        by = b_start_y + row * Inches(1.2)
        bcard = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, bx, by, Inches(2.15), Inches(1.05))
        bcard.fill.solid()
        bcard.fill.fore_color.rgb = c_navy if idx % 2 == 0 else c_subnavy
        bcard.line.fill.background()
        btf = bcard.text_frame
        btf.word_wrap = True
        p1 = btf.paragraphs[0]
        p1.text = title
        p1.font.bold = True
        p1.font.size = Pt(11)
        p1.font.color.rgb = c_gold
        p1.alignment = PP_ALIGN.CENTER
        p2 = btf.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(9.5)
        p2.font.color.rgb = c_white
        p2.alignment = PP_ALIGN.CENTER

    # Right Side Bottom: Primary Beneficiaries
    add_card(s5, Inches(7.8), Inches(5.85), Inches(4.733), Inches(1.25), "Primary beneficiaries")
    tb_ben = s5.shapes.add_textbox(Inches(7.95), Inches(6.25), Inches(4.433), Inches(0.75))
    tf_ben = tb_ben.text_frame
    tf_ben.word_wrap = True
    p_ben = tf_ben.paragraphs[0]
    p_ben.text = "Primary beneficiaries: warehouse operators, fulfillment centers, logistics teams and organizations deploying multi-AMR fleets."
    p_ben.font.size = Pt(11)
    p_ben.font.color.rgb = c_text_dark


    # ==========================================
    # SLIDE 6: RESEARCH AND REFERENCES
    # ==========================================
    s6 = prs.slides.add_slide(blank_slide_layout)
    set_bg(s6)
    add_header(s6, "RESEARCH AND REFERENCES", 6)

    # Left Side: IEEE / Research Foundations
    add_card(s6, Inches(0.8), Inches(1.3), Inches(6.8), Inches(5.3), "IEEE / Research Foundations")
    tb_ieee = s6.shapes.add_textbox(Inches(0.95), Inches(1.8), Inches(6.5), Inches(4.6))
    tf_ieee = tb_ieee.text_frame
    tf_ieee.word_wrap = True

    ieee_papers = [
        ("1. DC-MRTA — Decentralized Multi-Robot Task Allocation and Navigation in Complex Environments",
         "A. Agrawal et al., IEEE/RSJ IROS 2022. Combines decentralized task allocation with ORCA-based navigation for warehouse robots. DOI: 10.1109/IROS47612.2022.9981353",
         "https://ieeexplore.ieee.org/document/9981353"),
        ("2. RTAW — An Attention Inspired Reinforcement Learning Method for Multi-Robot Task Allocation in Warehouse Environments",
         "IEEE ICRA 2023. Studies multi-robot task allocation in warehouse environments and reports large-scale simulation results. DOI: 10.1109/ICRA48891.2023.10161310",
         "https://ieeexplore.ieee.org/document/10161310"),
        ("3. A Multi-robot Task Allocation and Path Planning Method for Warehouse System",
         "IEEE conference publication on joint multi-robot task allocation and path planning for warehouse systems.",
         "https://ieeexplore.ieee.org/document/9549796"),
        ("4. Decentralized Task Allocation for Redundant Multi-Robot Systems: An Iterative Consensus Approach",
         "IEEE ICCA 2024. Uses local information for decentralized task allocation and considers fault resilience.",
         "https://ieeexplore.ieee.org/document/10591823")
    ]
    for title, detail, url in ieee_papers:
        p1 = tf_ieee.add_paragraph()
        r1 = p1.add_run()
        r1.text = title
        r1.font.bold = True
        r1.font.size = Pt(10)
        r1.font.color.rgb = c_subnavy

        p2 = tf_ieee.add_paragraph()
        r2 = p2.add_run()
        r2.text = detail
        r2.font.size = Pt(9.5)
        r2.font.color.rgb = c_text_dark

        p3 = tf_ieee.add_paragraph()
        r3 = p3.add_run()
        r3.text = url
        r3.font.size = Pt(9)
        r3.font.color.rgb = c_blue_accent
        r3.hyperlink.address = url

    # Right Side: WAREHOUSE VIDEO REFERENCE
    add_card(s6, Inches(7.8), Inches(1.3), Inches(4.733), Inches(5.3), "WAREHOUSE VIDEO REFERENCE")
    tb_vid = s6.shapes.add_textbox(Inches(7.95), Inches(1.8), Inches(4.433), Inches(4.6))
    tf_vid = tb_vid.text_frame
    tf_vid.word_wrap = True

    p_v1 = tf_vid.paragraphs[0]
    p_v1.text = "World's most advanced robotic warehouse (AI automation)"
    p_v1.font.bold = True
    p_v1.font.size = Pt(12)
    p_v1.font.color.rgb = c_subnavy

    p_v2 = tf_vid.add_paragraph()
    p_v2.text = "Brightpick • 2024"
    p_v2.font.bold = True
    p_v2.font.size = Pt(10.5)
    p_v2.font.color.rgb = c_gold

    p_v3 = tf_vid.add_paragraph()
    p_v3.text = "A useful visual reference for ecommerce warehouse AMRs, robotic picking, navigation, order buffering, packout and software coordination."
    p_v3.font.size = Pt(10)
    p_v3.font.color.rgb = c_text_dark

    p_v4 = tf_vid.add_paragraph()
    p_v4.text = "YouTube video:\nhttps://www.youtube.com/watch?v=U2AGLeJBFNg"
    p_v4.font.size = Pt(10)
    p_v4.font.color.rgb = c_blue_accent

    p_v5 = tf_vid.add_paragraph()
    p_v5.text = "The video had 867K+ views at the time of research and shows a real ecommerce-warehouse workflow that is useful as a visual reference for our simulation."
    p_v5.font.size = Pt(9.5)
    p_v5.font.color.rgb = c_text_muted

    p_v6 = tf_vid.add_paragraph()
    p_v6.text = "Project Repository & Documentation:"
    p_v6.font.bold = True
    p_v6.font.size = Pt(10.5)
    p_v6.font.color.rgb = c_subnavy

    p_v7 = tf_vid.add_paragraph()
    r_doc = p_v7.add_run()
    r_doc.text = "https://github.com/karthik-ak-Git/NEXUS-Fleet"
    r_doc.font.size = Pt(9.5)
    r_doc.font.bold = True
    r_doc.font.color.rgb = c_blue_accent
    r_doc.hyperlink.address = "https://github.com/karthik-ak-Git/NEXUS-Fleet"

    # Footer Banner across bottom
    ft_bar = s6.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Inches(6.8), Inches(13.333), Inches(0.7))
    ft_bar.fill.solid()
    ft_bar.fill.fore_color.rgb = c_navy
    ft_bar.line.fill.background()

    ft_tb = s6.shapes.add_textbox(Inches(0.8), Inches(6.85), Inches(11.733), Inches(0.6))
    p_ft = ft_tb.text_frame.paragraphs[0]
    p_ft.text = "NEXUS-Fleet • SIH26123 • SPARK-08"
    p_ft.font.bold = True
    p_ft.font.size = Pt(14)
    p_ft.font.color.rgb = c_white
    p_ft.alignment = PP_ALIGN.CENTER

    # Save outputs
    out_file1 = r"C:\Users\Atina\Downloads\SIH_2026_Spark-08_Redesigned.pptx"
    out_file2 = r"d:\NEXUS-Fleet-Warehouse-Simulation\SIH_2026_Spark-08_Redesigned.pptx"

    prs.save(out_file1)
    prs.save(out_file2)
    print("Successfully generated redesigned presentation at:")
    print(" -", out_file1)
    print(" -", out_file2)

if __name__ == "__main__":
    create_deck()
